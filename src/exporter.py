from pathlib import Path
from typing import Sequence

from openpyxl import Workbook
from openpyxl.styles import (
    Font,
    PatternFill,
    Alignment,
)
from openpyxl.utils import get_column_letter

from .config import MiningConfig
from .models import SamplingRecord


class ExcelExporter:
    def __init__(self, config: MiningConfig):
        self.config = config

    def export(
        self,
        records: Sequence[SamplingRecord],
        output_file: Path,
        source_file: Path,
    ) -> None:

        output_file = Path(output_file)

        output_file.parent.mkdir(
            parents=True,
            exist_ok=True,
        )

        workbook = Workbook()

        worksheet = workbook.active
        worksheet.title = (
            self.config.output_sheet_name
        )

        self._write_data(
            worksheet,
            records,
        )

        self._write_control_sheet(
            workbook,
            records,
            source_file,
        )

        workbook.save(output_file)

    # -----------------------------------------------------

    def _write_data(
        self,
        worksheet,
        records: Sequence[SamplingRecord],
    ) -> None:

        worksheet.append(
            list(self.config.output_headers)
        )

        for record in records:
            worksheet.append(record.to_row())

        self._format_header(worksheet)

        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = (
            worksheet.dimensions
        )

        widths = {
            "A": 13,
            "B": 10,
            "C": 40,
            "D": 22,
            "E": 12,
            "F": 12,
            "G": 12,
            "H": 15,
        }

        for column, width in widths.items():
            worksheet.column_dimensions[
                column
            ].width = width

        for cell in worksheet["A"][1:]:
            cell.number_format = "dd/mm/yyyy"

    # -----------------------------------------------------

    @staticmethod
    def _format_header(worksheet) -> None:

        fill = PatternFill(
            fill_type="solid",
            fgColor="1F4E78",
        )

        font = Font(
            color="FFFFFF",
            bold=True,
        )

        alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

        for cell in worksheet[1]:

            cell.fill = fill
            cell.font = font
            cell.alignment = alignment

    # -----------------------------------------------------

    def _write_control_sheet(
        self,
        workbook,
        records,
        source_file,
    ) -> None:

        worksheet = workbook.create_sheet(
            self.config.control_sheet_name
        )

        controls = [
            ("Item", "Keterangan"),
            (
                "Source file",
                source_file.name,
            ),
            (
                "Source worksheet",
                self.config.worksheet_name,
            ),
            (
                "Output records",
                len(records),
            ),
            (
                "OFF handling",
                "Tidak dibuat sebagai record pengukuran",
            ),
            (
                "Total handling",
                "Bukan parameter; menjadi batas pembacaan SP",
            ),
            (
                "Blank value",
                "Tidak dibuat sebagai record pengukuran",
            ),
            (
                "STD '-'",
                "Dipertahankan sesuai source; tidak dianggap 0",
            ),
        ]

        for row in controls:
            worksheet.append(row)

        self._format_header(worksheet)

        worksheet.column_dimensions["A"].width = 24
        worksheet.column_dimensions["B"].width = 65