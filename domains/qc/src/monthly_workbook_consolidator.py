from __future__ import annotations

from collections import Counter
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from .monthly_source_planner import MonthlyDiscoveryCondition, MonthlyWorkItem


PROVENANCE_HEADERS = [
    "Source Month",
    "Source Week",
    "Domain",
    "Source Workbook",
]


class MonthlyWorkbookConsolidator:
    """Build the consolidated monthly workbook from per-workbook outputs."""

    def __init__(self):
        self.workbook = Workbook()
        self.ws_data = self.workbook.active
        self.ws_data.title = "Data_Mining"
        self.ws_control = self.workbook.create_sheet("Control")
        self.ws_log = self.workbook.create_sheet("Processing_Log")
        self.ws_summary = self.workbook.create_sheet("Run_Summary")

        self.data_header_state = {"header": None}
        self.log_header_state = {"header": None}
        self.totals = Counter()

        self.ws_summary.append([
            "Source Month",
            "Source Week",
            "Domain",
            "Source Workbook",
            "Source Path",
            "Execution",
            "Candidate Records",
            "Worksheets",
            "Success",
            "Excluded",
            "Error",
            "Message",
        ])

    def record_discovery_condition(
        self,
        condition: MonthlyDiscoveryCondition,
    ) -> None:
        source_path = "" if condition.source_path is None else str(condition.source_path)

        self.ws_summary.append([
            condition.month,
            condition.week,
            condition.domain,
            "",
            source_path,
            condition.status,
            0, 0, 0, 0, 0,
            condition.message,
        ])

        if condition.status == "MISSING_DOMAIN_FOLDER":
            self.totals["missing_domain_folder"] += 1
        elif condition.status == "NO_WORKBOOKS":
            self.totals["no_workbooks"] += 1

    def record_execution_failure(
        self,
        item: MonthlyWorkItem,
        message: str,
    ) -> None:
        self.totals["workbooks_attempted"] += 1
        self.totals["exec_failed"] += 1

        self.ws_summary.append([
            *item.provenance,
            str(item.source_path),
            "EXEC_FAILED",
            0, 0, 0, 0, 0,
            message,
        ])

    def consolidate_workbook(
        self,
        item: MonthlyWorkItem,
        temp_output: Path,
    ) -> None:
        self.totals["workbooks_attempted"] += 1

        execution = "OK"
        message = ""
        candidate_records = 0
        counts = Counter()

        try:
            wb_src = load_workbook(
                temp_output,
                data_only=True,
                read_only=True,
            )
            try:
                required = {"Data_Mining", "Control", "Processing_Log"}
                missing_sheets = required - set(wb_src.sheetnames)
                if missing_sheets:
                    raise ValueError(
                        "Missing output sheet(s): "
                        + ", ".join(sorted(missing_sheets))
                    )

                data_rows = self._read_sheet_rows(wb_src["Data_Mining"])
                control_rows = self._read_sheet_rows(wb_src["Control"])
                log_rows = self._read_sheet_rows(wb_src["Processing_Log"])

                candidate_records = self._append_with_provenance(
                    self.ws_data,
                    item.provenance,
                    data_rows,
                    self.data_header_state,
                    "Data_Mining",
                )

                self._append_with_provenance(
                    self.ws_log,
                    item.provenance,
                    log_rows,
                    self.log_header_state,
                    "Processing_Log",
                )

                self._add_control_section(
                    item.provenance,
                    control_rows,
                )

                counts = self._processing_counts(log_rows)
            finally:
                wb_src.close()

        except Exception as exc:
            execution = "OUTPUT_READ_ERROR"
            message = f"{type(exc).__name__}: {exc}"
            self.totals["output_read_error"] += 1

        worksheets = (
            counts.get("SUCCESS", 0)
            + counts.get("EXCLUDED", 0)
            + counts.get("ERROR", 0)
        )

        self.totals["candidate_records"] += candidate_records
        self.totals["worksheets"] += worksheets
        self.totals["success"] += counts.get("SUCCESS", 0)
        self.totals["excluded"] += counts.get("EXCLUDED", 0)
        self.totals["error"] += counts.get("ERROR", 0)

        self.ws_summary.append([
            *item.provenance,
            str(item.source_path),
            execution,
            candidate_records,
            worksheets,
            counts.get("SUCCESS", 0),
            counts.get("EXCLUDED", 0),
            counts.get("ERROR", 0),
            message,
        ])

    def finalize(self, final_output: Path) -> Counter:
        if self.data_header_state["header"] is None:
            self.ws_data.append(PROVENANCE_HEADERS + [
                "Source Worksheet",
                "Tanggal",
                "Jam",
                "Sampling Point (SP)",
                "Parameter",
                "Unit",
                "Min",
                "Max",
                "Nilai",
            ])

        if self.log_header_state["header"] is None:
            self.ws_log.append(PROVENANCE_HEADERS + [
                "Worksheet",
                "Visibility",
                "Status",
                "Record Count",
                "Message",
            ])

        self._set_simple_formatting(self.ws_data, freeze="A2", autofilter=True)
        self._set_simple_formatting(self.ws_log, freeze="A2", autofilter=True)
        self._set_simple_formatting(self.ws_summary, freeze="A2", autofilter=True)
        self._set_simple_formatting(self.ws_control, freeze=None, autofilter=False)

        final_output = Path(final_output)
        if final_output.exists():
            final_output.unlink()

        self.workbook.save(final_output)
        return Counter(self.totals)

    @staticmethod
    def _read_sheet_rows(ws):
        return [list(row) for row in ws.iter_rows(values_only=True)]

    @staticmethod
    def _normalize_header(row) -> list[str]:
        return ["" if value is None else str(value).strip() for value in row]

    @classmethod
    def _append_with_provenance(
        cls,
        ws_out,
        provenance,
        rows,
        expected_header_state,
        sheet_name,
    ) -> int:
        if not rows:
            return 0

        header = cls._normalize_header(rows[0])

        if expected_header_state["header"] is None:
            expected_header_state["header"] = header
            ws_out.append(PROVENANCE_HEADERS + header)
        elif header != expected_header_state["header"]:
            raise ValueError(
                f"{sheet_name} header mismatch. "
                f"Expected={expected_header_state['header']!r}; Actual={header!r}"
            )

        count = 0
        for row in rows[1:]:
            if not any(
                value is not None and str(value).strip() != ""
                for value in row
            ):
                continue
            ws_out.append(list(provenance) + list(row))
            count += 1

        return count

    @classmethod
    def _processing_counts(cls, rows) -> Counter:
        result = Counter()
        if not rows:
            return result

        header = cls._normalize_header(rows[0])
        index = {name.casefold(): i for i, name in enumerate(header)}
        status_idx = index.get("status")
        if status_idx is None:
            return result

        for row in rows[1:]:
            if status_idx >= len(row):
                continue
            value = row[status_idx]
            if value is None:
                continue
            result[str(value).strip().upper()] += 1

        return result

    def _add_control_section(self, provenance, control_rows) -> None:
        self.ws_control.append([])
        self.ws_control.append([
            "Source Month", provenance[0],
            "Source Week", provenance[1],
            "Domain", provenance[2],
            "Source Workbook", provenance[3],
        ])

        if not control_rows:
            self.ws_control.append(["(Control sheet empty)"])
            return

        for row in control_rows:
            self.ws_control.append(list(row))

    @staticmethod
    def _set_simple_formatting(
        ws,
        freeze="A2",
        autofilter=True,
        max_width=45,
    ) -> None:
        if ws.max_row >= 1:
            for cell in ws[1]:
                cell.font = Font(bold=True)

        if freeze:
            ws.freeze_panes = freeze

        if autofilter and ws.max_row >= 1 and ws.max_column >= 1:
            ws.auto_filter.ref = ws.dimensions

        sample_rows = min(ws.max_row, 250)
        for col_idx in range(1, ws.max_column + 1):
            width = 10
            for row_idx in range(1, sample_rows + 1):
                value = ws.cell(row=row_idx, column=col_idx).value
                if value is None:
                    continue
                width = max(width, min(len(str(value)) + 2, max_width))
            ws.column_dimensions[get_column_letter(col_idx)].width = width
