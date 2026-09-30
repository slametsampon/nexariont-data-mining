from dataclasses import dataclass
from pathlib import Path
from typing import Any

import xlrd
from openpyxl import load_workbook


@dataclass(frozen=True)
class WorksheetInfo:
    """Metadata worksheet untuk penemuan dan penyaringan berdasarkan visibilitas.

    Attributes:
        name (str): Nama worksheet; argumen konstruktor dataclass.
        visibility (int): Status sumber: 0 visible, 1 hidden, 2 very hidden,
            atau -1 untuk status .xlsx yang tidak dikenali.
        is_visible (bool): Properti bernilai True hanya jika visibility adalah 0.

    Examples:
        >>> WorksheetInfo("shift-pagi", 0).is_visible
        True
    """

    name: str
    visibility: int

    @property
    def is_visible(self) -> bool:
        return self.visibility == 0


class CellAdapter:
    """Representasi sel minimum yang dibutuhkan parser laporan QC.

    Args:
        value (Any): Nilai sel yang telah dikonversi dari xlrd.
        row (int): Nomor baris bergaya Excel, dimulai dari 1.
        column (int): Nomor kolom bergaya Excel, dimulai dari 1.

    Attributes:
        value (Any): Isi sel; None untuk sel kosong atau di luar batas sheet.
        row (int): Nomor baris yang diminta.
        column (int): Nomor kolom yang diminta.
    """

    def __init__(
        self,
        value: Any,
        row: int,
        column: int,
    ):
        self.value = value
        self.row = row
        self.column = column


class XlsWorksheetAdapter:
    """Menyediakan antarmuka worksheet bagi sumber .xls yang dibaca xlrd.

    Menyediakan cell() dan iter_rows() dengan indeks mulai dari 1 agar parser
    dapat membaca .xls tanpa membuat file .xlsx sementara. Nilai tanggal serial
    xlrd tidak dikonversi menjadi objek tanggal oleh adapter ini.

    Args:
        sheet (xlrd.sheet.Sheet): Worksheet sumber dari workbook xlrd.

    Attributes:
        max_row (int): Jumlah baris pada worksheet sumber.
        max_column (int): Jumlah kolom pada worksheet sumber.
    """

    def __init__(self, sheet: xlrd.sheet.Sheet):
        self._sheet = sheet

        self.max_row = sheet.nrows
        self.max_column = sheet.ncols

    def cell(
        self,
        row: int,
        column: int,
    ) -> CellAdapter:

        # Parser menggunakan index Excel mulai dari 1,
        # sedangkan xlrd menggunakan index mulai dari 0.
        row_index = row - 1
        col_index = column - 1

        if (
            row_index < 0
            or col_index < 0
            or row_index >= self._sheet.nrows
            or col_index >= self._sheet.ncols
        ):
            return CellAdapter(
                value=None,
                row=row,
                column=column,
            )

        xls_cell = self._sheet.cell(
            row_index,
            col_index,
        )

        value = self._convert_value(
            xls_cell
        )

        return CellAdapter(
            value=value,
            row=row,
            column=column,
        )

    def iter_rows(
        self,
        min_row: int = 1,
        max_row: int | None = None,
        min_col: int = 1,
        max_col: int | None = None,
    ):

        if max_row is None:
            max_row = self.max_row

        if max_col is None:
            max_col = self.max_column

        for row_number in range(
            min_row,
            max_row + 1,
        ):

            row_cells = []

            for column_number in range(
                min_col,
                max_col + 1,
            ):

                row_cells.append(
                    self.cell(
                        row=row_number,
                        column=column_number,
                    )
                )

            yield tuple(row_cells)

    @staticmethod
    def _convert_value(
        cell: xlrd.sheet.Cell,
    ) -> Any:
        """
        Konversi minimum dan konservatif.

        Blank -> None

        Text seperti:
            <0.01
            -
            OFF
        tetap dipertahankan sebagai text.

        Numeric tetap numeric.
        """

        if cell.ctype in (
            xlrd.XL_CELL_EMPTY,
            xlrd.XL_CELL_BLANK,
        ):
            return None

        if cell.ctype == xlrd.XL_CELL_TEXT:
            return cell.value

        if cell.ctype == xlrd.XL_CELL_NUMBER:

            value = cell.value

            # 5.0 -> 5
            # 9.2282 -> 9.2282
            if float(value).is_integer():
                return int(value)

            return value

        if cell.ctype == xlrd.XL_CELL_BOOLEAN:
            return bool(cell.value)

        if cell.ctype == xlrd.XL_CELL_ERROR:
            return cell.value

        return cell.value


class WorkbookReader:
    """Membaca worksheet dan metadata workbook .xls atau .xlsx.

    Format .xls menggunakan xlrd dan XlsWorksheetAdapter. Format .xlsx
    menggunakan openpyxl dengan data_only=True dan read_only=True, sehingga
    nilai formula mengikuti cache workbook. Konstruktor hanya menyimpan path;
    validasi file dilakukan saat metode pembacaan dipanggil.

    Args:
        file_path (Path): Lokasi workbook sumber.

    Attributes:
        file_path (Path): Path sumber yang telah dikonversi menjadi Path.

    Examples:
        reader = WorkbookReader(Path("laporan.xlsx"))
        names = reader.get_visible_worksheet_names()
        worksheet = reader.load_worksheet(names[0])
    """

    def __init__(
        self,
        file_path: Path,
    ):
        self.file_path = Path(file_path)

    def get_worksheet_names(
        self,
    ) -> list[str]:
        """
        Mengembalikan seluruh nama worksheet dari workbook.

        Supported:
        - .xls  : xlrd
        - .xlsx : openpyxl

        Method ini hanya melakukan worksheet discovery.
        Tidak melakukan parsing data.
        """

        if not self.file_path.exists():
            raise FileNotFoundError(
                f"File tidak ditemukan: {self.file_path}"
            )

        extension = self.file_path.suffix.lower()

        if extension == ".xls":

            workbook = xlrd.open_workbook(
                filename=str(self.file_path),
                on_demand=True,
            )

            try:
                return workbook.sheet_names()

            finally:
                workbook.release_resources()

        if extension == ".xlsx":

            workbook = load_workbook(
                filename=self.file_path,
                data_only=True,
                read_only=True,
            )

            try:
                return list(
                    workbook.sheetnames
                )

            finally:
                workbook.close()

        raise ValueError(
            "Format input tidak didukung. "
            "Gunakan file .xls atau .xlsx"
        )

    def get_worksheet_info(
        self,
    ) -> list[WorksheetInfo]:
        """
        Mengembalikan metadata seluruh worksheet,
        termasuk visibility.

        Visible worksheet dapat menjadi candidate parsing.
        Hidden/non-visible worksheet tidak menjadi candidate.
        """

        if not self.file_path.exists():
            raise FileNotFoundError(
                f"File tidak ditemukan: {self.file_path}"
            )

        extension = self.file_path.suffix.lower()

        if extension == ".xls":

            workbook = xlrd.open_workbook(
                filename=str(self.file_path),
                on_demand=True,
            )

            try:
                return [
                    WorksheetInfo(
                        name=sheet.name,
                        visibility=sheet.visibility,
                    )
                    for sheet in workbook.sheets()
                ]

            finally:
                workbook.release_resources()

        if extension == ".xlsx":

            workbook = load_workbook(
                filename=self.file_path,
                data_only=True,
                read_only=True,
            )

            try:
                visibility_map = {
                    "visible": 0,
                    "hidden": 1,
                    "veryHidden": 2,
                }

                return [
                    WorksheetInfo(
                        name=sheet.title,
                        visibility=visibility_map.get(
                            sheet.sheet_state,
                            -1,
                        ),
                    )
                    for sheet in workbook.worksheets
                ]

            finally:
                workbook.close()

        raise ValueError(
            "Format input tidak didukung. "
            "Gunakan file .xls atau .xlsx"
        )

    def get_visible_worksheet_names(
        self,
    ) -> list[str]:
        """
        Mengembalikan hanya worksheet yang visible.

        Hidden/non-visible worksheet tidak menjadi
        candidate parsing.
        """

        return [
            worksheet.name
            for worksheet in self.get_worksheet_info()
            if worksheet.is_visible
        ]

    def load_worksheet(
        self,
        worksheet_name: str,
    ):

        """Membaca satu worksheet melalui reader yang sesuai ekstensi sumber.

        Args:
            worksheet_name (str): Nama worksheet persis seperti di workbook.

        Returns:
            XlsWorksheetAdapter | openpyxl.worksheet._read_only.ReadOnlyWorksheet:
            Worksheet dengan antarmuka cell(), iter_rows(), max_row, dan max_column.

        Raises:
            FileNotFoundError: File sumber tidak ditemukan.
            ValueError: Ekstensi tidak didukung atau nama worksheet tidak ditemukan.
        """

        if not self.file_path.exists():
            raise FileNotFoundError(
                f"File tidak ditemukan: {self.file_path}"
            )

        extension = self.file_path.suffix.lower()

        if extension == ".xls":
            return self._load_xls(
                worksheet_name
            )

        if extension == ".xlsx":
            return self._load_xlsx(
                worksheet_name
            )

        raise ValueError(
            "Format input tidak didukung. "
            "Gunakan file .xls atau .xlsx"
        )

    def _load_xls(
        self,
        worksheet_name: str,
    ) -> XlsWorksheetAdapter:

        workbook = xlrd.open_workbook(
            filename=str(self.file_path),
            on_demand=True,
        )

        if worksheet_name not in workbook.sheet_names():

            available = ", ".join(
                workbook.sheet_names()
            )

            raise ValueError(
                f"Worksheet '{worksheet_name}' "
                f"tidak ditemukan. "
                f"Worksheet tersedia: {available}"
            )

        sheet = workbook.sheet_by_name(
            worksheet_name
        )

        return XlsWorksheetAdapter(
            sheet
        )

    def _load_xlsx(
        self,
        worksheet_name: str,
    ):

        workbook = load_workbook(
            filename=self.file_path,
            data_only=True,
            read_only=True,
        )

        if worksheet_name not in workbook.sheetnames:

            available = ", ".join(
                workbook.sheetnames
            )

            raise ValueError(
                f"Worksheet '{worksheet_name}' "
                f"tidak ditemukan. "
                f"Worksheet tersedia: {available}"
            )

        return workbook[
            worksheet_name
        ]