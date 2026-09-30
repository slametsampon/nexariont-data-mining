import re
from dataclasses import dataclass
from pathlib import Path
from typing import Optional

from openpyxl import load_workbook


@dataclass(frozen=True)
class SamplingPointMasterEntry:
    """Satu entri referensi titik sampling dari worksheet master.

    Dataclass frozen ini merekam tiga kolom pertama master dan nomor baris
    asalnya. Semua atribut merupakan argumen konstruktor.

    Attributes:
        domain (str): Domain proses dari kolom A.
        sampling_identity (str): Identitas utama titik sampling dari kolom B.
        description (str | None): Keterangan kolom C, atau None jika kosong.
        row_number (int): Nomor baris Excel sumber, dimulai dari 1.
    """

    domain: str
    sampling_identity: str
    description: Optional[str]
    row_number: int


class SamplingPointMaster:
    """Referensi terkontrol untuk pengenalan identitas titik sampling.

    Pencocokan dibatasi pada identitas kolom B yang dinormalisasi, kombinasi
    B+C yang didukung, atau identitas dengan akhiran waktu (HH.MM)/(HH:MM).
    Normalisasi mengabaikan kapitalisasi dan karakter selain a-z serta 0-9.
    Hasil ambigu dikembalikan sebagai None; tidak dilakukan pencocokan fuzzy.

    Args:
        entries (list[SamplingPointMasterEntry]): Entri yang diindeks per domain.

    Attributes:
        entries (list[SamplingPointMasterEntry]): Daftar entri yang diberikan.
            Hindari mutasi setelah konstruksi karena indeks domain dibangun sekali.

    Examples:
        master = SamplingPointMaster.load(Path("master.xlsx"))
        entry = master.match("NPG", "SP-01")
    """

    def __init__(self, entries: list[SamplingPointMasterEntry]):
        self.entries = entries
        self._by_domain: dict[str, list[SamplingPointMasterEntry]] = {}
        for entry in entries:
            self._by_domain.setdefault(self._norm(entry.domain), []).append(entry)

    @classmethod
    def load(cls, path: Path, worksheet_name: str = "QA Review") -> "SamplingPointMaster":
        """Membaca entri master dari tiga kolom pertama, mulai baris kedua.

        Args:
            path (Path): Lokasi workbook master yang didukung openpyxl.
            worksheet_name (str): Nama worksheet master; default "QA Review".

        Returns:
            SamplingPointMaster: Referensi dengan indeks domain. Baris tanpa domain
            atau identitas dilewati; deskripsi kosong atau "(blank)" menjadi None.

        Raises:
            FileNotFoundError: File master tidak ditemukan.
            ValueError: Worksheet master tidak ditemukan.
        """

        path = Path(path)
        if not path.exists():
            raise FileNotFoundError(f"Master SSP tidak ditemukan: {path}")
        wb = load_workbook(path, data_only=True, read_only=True)
        try:
            if worksheet_name not in wb.sheetnames:
                raise ValueError(f"Worksheet master SSP '{worksheet_name}' tidak ditemukan")
            ws = wb[worksheet_name]
            entries = []
            for row_number, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                domain, identity, description = (list(row) + [None, None, None])[:3]
                if cls._blank(domain) or cls._blank(identity):
                    continue
                desc = None if cls._blank(description) or str(description).strip().casefold() == "(blank)" else str(description).strip()
                entries.append(SamplingPointMasterEntry(str(domain).strip(), str(identity).strip(), desc, row_number))
            return cls(entries)
        finally:
            wb.close()

    def match(self, domain: str, source_text) -> Optional[SamplingPointMasterEntry]:
        """Mencari satu entri unik dalam domain berdasarkan teks sumber.

        Args:
            domain (str): Domain yang membatasi kandidat pencocokan.
            source_text (Any): Nilai sel berisi identitas titik sampling.

        Returns:
            SamplingPointMasterEntry | None: Entri unik, atau None jika teks kosong,
            tidak cocok, atau ambigu.
        """

        if self._blank(source_text):
            return None
        candidates = self._by_domain.get(self._norm(domain), [])
        source = str(source_text).strip()
        source_key = self._norm(source)

        # 1. Exact primary SSP identity.
        primary = [e for e in candidates if self._norm(e.sampling_identity) == source_key]
        if len(primary) == 1:
            return primary[0]
        if len(primary) > 1:
            return None  # B alone is ambiguous; B+C is required.

        # 2. Exact controlled B+C textual forms.
        combined = []
        for e in candidates:
            if not e.description:
                continue
            forms = (
                f"{e.sampling_identity} ({e.description})",
                f"{e.sampling_identity} {e.description}",
                f"{e.sampling_identity}|{e.description}",
            )
            if source_key in {self._norm(x) for x in forms}:
                combined.append(e)
        if len(combined) == 1:
            return combined[0]
        if len(combined) > 1:
            return None

        # 3. User-established P-xxx case: remove trailing time only when the
        # stripped base is itself an exact, unique master SSP in this domain.
        m = re.fullmatch(r"\s*(.*?)\s*\(\s*\d{1,2}[.:]\d{2}\s*\)\s*", source)
        if m:
            base_key = self._norm(m.group(1))
            stripped = [e for e in candidates if self._norm(e.sampling_identity) == base_key]
            if len(stripped) == 1:
                return stripped[0]

        return None

    def output_identity(
        self,
        domain: str,
        source_text,
        matched_entry: SamplingPointMasterEntry,
    ) -> str:
        """Return output identity without collapsing an ambiguous primary B.

        A unique primary B keeps the existing normalized Master identity.
        When the same B exists more than once in the domain, a successful B+C
        match keeps the exact source identity so distinct Sampling Points remain
        distinct without inventing a new canonical B+C representation.
        """
        candidates = self._by_domain.get(self._norm(domain), [])
        primary_key = self._norm(matched_entry.sampling_identity)
        same_primary = [
            entry
            for entry in candidates
            if self._norm(entry.sampling_identity) == primary_key
        ]

        if len(same_primary) > 1:
            return str(source_text).strip()

        return matched_entry.sampling_identity

    @staticmethod
    def _norm(value) -> str:
        if value is None:
            return ""
        return re.sub(r"[^a-z0-9]+", "", str(value).casefold())

    @staticmethod
    def _blank(value) -> bool:
        return value is None or (isinstance(value, str) and not value.strip())
