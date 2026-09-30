from __future__ import annotations

import tempfile
from collections import Counter
from dataclasses import dataclass
from pathlib import Path
from typing import Callable

from openpyxl import Workbook, load_workbook

from .monthly_source_planner import MonthlySourcePlanner
from .monthly_workbook_consolidator import MonthlyWorkbookConsolidator
from .workbook_processor import QCWorkbookProcessor


MASTER_SOURCE_SHEET = "Sampling-Point"
CORE_MASTER_SHEET = "QA Review"

ProgressCallback = Callable[[str], None]


class NoWeekFoldersFoundError(ValueError):
    """Menandakan tidak ada folder minggu dalam rencana pemrosesan bulanan.

    Subclass ValueError yang dilempar MonthlyMiningOrchestrator.run() sebelum
    pemrosesan master dan workbook dimulai.

    Args:
        *args: Argumen exception bawaan; biasanya pesan berisi folder input.
    """

    pass


@dataclass(frozen=True)
class MonthlyRunResult:
    """Ringkasan keluaran satu proses mining QC bulanan.

    Seluruh atribut merupakan argumen konstruktor dataclass frozen. totals
    tetap merupakan Counter mutable meskipun atribut dataclass bersifat frozen.

    Attributes:
        input_root (Path): Folder bulan sumber.
        output_dir (Path): Direktori keluaran yang diberikan pemanggil;
            disimpan sebagai metadata, bukan penentu lokasi final_output.
        master (Path): Path workbook master asli.
        final_output (Path): Path workbook konsolidasi yang telah disimpan.
        month (str): Nama folder bulan dari rencana.
        weeks_detected (int): Jumlah folder minggu yang ditemukan.
        totals (Counter): Hitungan candidate_records, worksheets, success,
            excluded, error, workbooks_attempted, exec_failed,
            output_read_error, missing_domain_folder, dan no_workbooks.
            Key yang tidak tercatat memiliki nilai baca 0.
    """

    input_root: Path
    output_dir: Path
    master: Path
    final_output: Path
    month: str
    weeks_detected: int
    totals: Counter


class MonthlyMiningOrchestrator:
    """Mengatur penemuan sumber, pemrosesan, dan konsolidasi QC bulanan.

    Worksheet master "Sampling-Point" disalin nilainya ke adapter sementara
    "QA Review". Setiap workbook diproses melalui QCWorkbookProcessor, lalu
    hasilnya digabungkan dengan informasi asal bulan, minggu, domain, dan file.
    Kegagalan pemrosesan workbook dicatat agar workbook berikutnya tetap diproses.

    Args:
        planner (MonthlySourcePlanner | None): Penyusun rencana sumber;
            default instance MonthlySourcePlanner.
        processor (QCWorkbookProcessor | None): Pemroses workbook;
            default instance QCWorkbookProcessor.

    Attributes:
        planner (MonthlySourcePlanner): Penyusun rencana yang digunakan.
        processor (QCWorkbookProcessor): Pemroses workbook yang digunakan.

    Examples:
        orchestrator = MonthlyMiningOrchestrator()
        result = orchestrator.run(
            input_root=Path("input/September 2026"),
            output_dir=Path("output"),
            master=Path("_Master-Data.xlsx"),
            final_output=Path("output/hasil_bulanan.xlsx"),
            progress=print,
        )
    """

    def __init__(
        self,
        planner: MonthlySourcePlanner | None = None,
        processor: QCWorkbookProcessor | None = None,
    ):
        self.planner = planner or MonthlySourcePlanner()
        self.processor = processor or QCWorkbookProcessor()

    def run(
        self,
        input_root: Path,
        output_dir: Path,
        master: Path,
        final_output: Path,
        progress: ProgressCallback | None = None,
    ) -> MonthlyRunResult:
        """Menjalankan rencana bulanan dan menyimpan workbook konsolidasi.

        Args:
            input_root (Path): Folder bulan berisi folder minggu.
            output_dir (Path): Metadata direktori keluaran untuk hasil eksekusi.
            master (Path): Workbook master dengan worksheet "Sampling-Point".
            final_output (Path): Lokasi hasil; direktori induk harus sudah tersedia.
                File yang sudah ada akan diganti oleh consolidator.
            progress (Callable[[str], None] | None): Callback pesan kemajuan;
                default None. Exception callback diteruskan kepada pemanggil.

        Returns:
            MonthlyRunResult: Lokasi hasil, jumlah minggu, dan hitungan pemrosesan.

        Raises:
            NoWeekFoldersFoundError: Tidak ditemukan folder minggu.
            FileNotFoundError: Folder input atau workbook master tidak tersedia.
            ValueError: Worksheet "Sampling-Point" tidak ditemukan pada master.
            OSError: Kegagalan akses berkas yang tidak ditangani per workbook.
        """

        input_root = Path(input_root)
        output_dir = Path(output_dir)
        master = Path(master)
        final_output = Path(final_output)

        plan = self.planner.plan(input_root)
        if not plan.weeks:
            raise NoWeekFoldersFoundError(
                f"NO_WEEK_FOLDERS_FOUND under: {input_root}"
            )

        consolidator = MonthlyWorkbookConsolidator()
        source_counter = 0
        current_week: str | None = None

        with tempfile.TemporaryDirectory(prefix="nexariont_qc_monthly_") as tmp:
            tmp_dir = Path(tmp)
            adapted_master = self._build_core_master_adapter(
                master,
                tmp_dir / "_Master-Data_core_adapter.xlsx",
            )

            for domain_plan in plan.domain_plans:
                if domain_plan.week != current_week:
                    current_week = domain_plan.week
                    self._emit(progress, f"\n=== {current_week} ===")

                if domain_plan.condition is not None:
                    if domain_plan.condition.status == "MISSING_DOMAIN_FOLDER":
                        self._emit(
                            progress,
                            f"[MISSING] {domain_plan.folder_name}",
                        )
                    elif domain_plan.condition.status == "NO_WORKBOOKS":
                        self._emit(
                            progress,
                            f"[EMPTY]   {domain_plan.folder_name}",
                        )

                    consolidator.record_discovery_condition(
                        domain_plan.condition
                    )
                    continue

                self._emit(
                    progress,
                    f"[{domain_plan.folder_name}] "
                    f"{len(domain_plan.work_items)} workbook(s)",
                )

                for item in domain_plan.work_items:
                    source_counter += 1
                    temp_output = tmp_dir / f"{source_counter:04d}.xlsx"

                    self._emit(
                        progress,
                        f"  {item.source_workbook}",
                    )

                    try:
                        self.processor.process(
                            input_file=item.source_path,
                            output_file=temp_output,
                            master_file=adapted_master,
                            domain=item.domain,
                        )
                    except Exception as exc:
                        consolidator.record_execution_failure(
                            item,
                            f"ERROR: {exc}",
                        )
                        continue

                    if not temp_output.exists():
                        consolidator.record_execution_failure(
                            item,
                            "Output file tidak terbentuk.",
                        )
                        continue

                    consolidator.consolidate_workbook(
                        item,
                        temp_output,
                    )

        totals = consolidator.finalize(final_output)

        return MonthlyRunResult(
            input_root=input_root,
            output_dir=output_dir,
            master=master,
            final_output=final_output,
            month=plan.month,
            weeks_detected=len(plan.weeks),
            totals=totals,
        )

    @staticmethod
    def _emit(
        progress: ProgressCallback | None,
        message: str,
    ) -> None:
        if progress is not None:
            progress(message)

    @staticmethod
    def _build_core_master_adapter(
        master_path: Path,
        adapter_path: Path,
    ) -> Path:
        wb_source = load_workbook(
            master_path,
            data_only=True,
            read_only=True,
        )
        try:
            if MASTER_SOURCE_SHEET not in wb_source.sheetnames:
                raise ValueError(
                    f"Worksheet master '{MASTER_SOURCE_SHEET}' tidak ditemukan "
                    f"dalam {master_path}"
                )

            ws_source = wb_source[MASTER_SOURCE_SHEET]

            wb_adapter = Workbook()
            ws_adapter = wb_adapter.active
            ws_adapter.title = CORE_MASTER_SHEET

            for row in ws_source.iter_rows(values_only=True):
                ws_adapter.append(list(row))

            wb_adapter.save(adapter_path)
            return adapter_path
        finally:
            wb_source.close()
