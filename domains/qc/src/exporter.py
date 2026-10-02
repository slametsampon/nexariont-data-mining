import ast
import re

from pathlib import Path
from typing import Sequence

from openpyxl import Workbook
from openpyxl.styles import (
    Font,
    PatternFill,
    Alignment,
)
from openpyxl.utils import get_column_letter

from .config import MiningConfig
from .models import (
    SamplingRecord,
    SourcedSamplingRecord,
)


class ExcelExporter:
    """Menulis candidate record QC ke workbook .xlsx beserta informasi kontrol.

    export() menulis delapan kolom data dan worksheet kontrol.
    export_multi_worksheet() menambahkan kolom Source Worksheet serta
    Processing_Log. Kedua metode membuat direktori induk bila diperlukan
    dan menyimpan workbook pada output_file, mengganti file yang sudah ada.

    Args:
        config (MiningConfig): Nama worksheet, header, dan konfigurasi keluaran.

    Attributes:
        config (MiningConfig): Konfigurasi ekspor yang digunakan.

    Examples:
        exporter = ExcelExporter(MiningConfig())
        exporter.export(
            records=[], output_file=Path("hasil.xlsx"),
            source_file=Path("laporan.xls"),
        )
    """

    def __init__(
        self,
        config: MiningConfig,
    ):
        self.config = config

    # =====================================================
    # EXISTING SINGLE-WORKSHEET EXPORT
    # =====================================================

    def export(
        self,
        records: Sequence[SamplingRecord],
        output_file: Path,
        source_file: Path,
    ) -> None:
        """Menyimpan record mode satu worksheet dan informasi kontrol.

        Args:
            records (Sequence[SamplingRecord]): Record yang akan diekspor.
            output_file (Path): Lokasi file .xlsx tujuan.
            source_file (Path): Path sumber untuk informasi penelusuran.

        Returns:
            None: Workbook disimpan ke output_file.
        """

        output_file = Path(
            output_file
        )

        output_file.parent.mkdir(
            parents=True,
            exist_ok=True,
        )

        workbook = Workbook()

        worksheet = workbook.active
        worksheet.title = (
            self.config.output_sheet_name
        )

        self._write_data(
            worksheet,
            records,
        )

        self._write_control_sheet(
            workbook,
            records,
            source_file,
        )

        workbook.save(
            output_file
        )

    # =====================================================
    # MULTI-WORKSHEET EXPORT — CP4.3
    # =====================================================

    def export_multi_worksheet(
        self,
        records: Sequence[SourcedSamplingRecord],
        worksheet_results,
        output_file: Path,
        source_file: Path,
    ) -> None:
        """Menyimpan data gabungan, kontrol, dan log pemrosesan worksheet.

        Candidate record merupakan hasil ekstraksi, bukan pernyataan validasi
        atau persetujuan QC. Nama worksheet asal disimpan pada setiap record.

        Args:
            records (Sequence[SourcedSamplingRecord]): Record beserta worksheet asal.
            worksheet_results (Iterable[WorksheetProcessingResult]): Ringkasan
                pemrosesan untuk ditulis pada Processing_Log.
            output_file (Path): Lokasi file .xlsx tujuan.
            source_file (Path): Path workbook sumber untuk informasi kontrol.

        Returns:
            None: Workbook disimpan ke output_file.
        """

        output_file = Path(
            output_file
        )

        source_file = Path(
            source_file
        )

        output_file.parent.mkdir(
            parents=True,
            exist_ok=True,
        )

        workbook = Workbook()

        worksheet = workbook.active
        worksheet.title = (
            self.config.output_sheet_name
        )

        self._write_multi_worksheet_data(
            worksheet=worksheet,
            records=records,
        )

        self._write_multi_worksheet_control_sheet(
            workbook=workbook,
            records=records,
            source_file=source_file,
        )

        self._write_processing_log(
            workbook=workbook,
            worksheet_results=worksheet_results,
        )

        workbook.save(
            output_file
        )

    # =====================================================
    # EXISTING SINGLE-WORKSHEET DATA WRITER
    # =====================================================

    def _write_data(
        self,
        worksheet,
        records: Sequence[SamplingRecord],
    ) -> None:

        worksheet.append(
            list(
                self.config.output_headers
            )
        )

        for record in records:
            worksheet.append(
                record.to_row()
            )

        self._format_header(
            worksheet
        )

        worksheet.freeze_panes = "A2"

        worksheet.auto_filter.ref = (
            worksheet.dimensions
        )

        widths = {
            "A": 13,
            "B": 10,
            "C": 40,
            "D": 22,
            "E": 12,
            "F": 12,
            "G": 12,
            "H": 15,
        }

        for column, width in (
            widths.items()
        ):
            worksheet.column_dimensions[
                column
            ].width = width

        for cell in worksheet["A"][1:]:
            cell.number_format = (
                "dd/mm/yyyy"
            )

    # =====================================================
    # MULTI-WORKSHEET DATA WRITER
    # =====================================================

    def _write_multi_worksheet_data(
        self,
        worksheet,
        records: Sequence[SourcedSamplingRecord],
    ) -> None:
        """
        Menulis consolidated multi-worksheet records.

        Struktur:
        Source Worksheet
        Tanggal
        Jam
        Sampling Point (SP)
        Parameter
        Unit
        Min
        Max
        Nilai
        """

        headers = [
            "Source Worksheet",
            *self.config.output_headers,
        ]

        worksheet.append(
            list(headers)
        )

        for sourced_record in records:
            worksheet.append(
                sourced_record.to_row()
            )

        self._format_header(
            worksheet
        )

        worksheet.freeze_panes = "A2"

        worksheet.auto_filter.ref = (
            worksheet.dimensions
        )

        widths = {
            "A": 24,
            "B": 13,
            "C": 10,
            "D": 40,
            "E": 22,
            "F": 12,
            "G": 12,
            "H": 12,
            "I": 15,
        }

        for column, width in (
            widths.items()
        ):
            worksheet.column_dimensions[
                column
            ].width = width

        # Pada multi-worksheet export:
        #
        # A = Source Worksheet
        # B = Tanggal
        #
        # Existing single-worksheet export tetap
        # menggunakan kolom A sebagai Tanggal.
        for cell in worksheet["B"][1:]:
            cell.number_format = (
                "dd/mm/yyyy"
            )

    # =====================================================
    # SHARED HEADER FORMAT
    # =====================================================

    @staticmethod
    def _format_header(
        worksheet,
    ) -> None:

        fill = PatternFill(
            fill_type="solid",
            fgColor="1F4E78",
        )

        font = Font(
            color="FFFFFF",
            bold=True,
        )

        alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

        for cell in worksheet[1]:
            cell.fill = fill
            cell.font = font
            cell.alignment = alignment

    # =====================================================
    # EXISTING SINGLE-WORKSHEET CONTROL
    # =====================================================

    def _write_control_sheet(
        self,
        workbook,
        records,
        source_file,
    ) -> None:

        worksheet = workbook.create_sheet(
            self.config.control_sheet_name
        )

        controls = [
            (
                "Item",
                "Keterangan",
            ),
            (
                "Source file",
                source_file.name,
            ),
            (
                "Source worksheet",
                self.config.worksheet_name,
            ),
            (
                "Output records",
                len(records),
            ),
            (
                "OFF handling",
                (
                    "Tidak dibuat sebagai "
                    "record pengukuran"
                ),
            ),
            (
                "Total handling",
                (
                    "Bukan parameter; menjadi "
                    "batas pembacaan SP"
                ),
            ),
            (
                "Blank value",
                (
                    "Tidak dibuat sebagai "
                    "record pengukuran"
                ),
            ),
            (
                "STD '-'",
                (
                    "Dipertahankan sesuai source; "
                    "tidak dianggap 0"
                ),
            ),
        ]

        for row in controls:
            worksheet.append(
                row
            )

        self._format_header(
            worksheet
        )

        worksheet.column_dimensions[
            "A"
        ].width = 24

        worksheet.column_dimensions[
            "B"
        ].width = 65

    # =====================================================
    # MULTI-WORKSHEET CONTROL — CP4.3
    # =====================================================

    def _write_multi_worksheet_control_sheet(
        self,
        workbook,
        records: Sequence[SourcedSamplingRecord],
        source_file: Path,
    ) -> None:
        """
        Control minimum untuk multi-worksheet export.

        Tidak menyatakan candidate records sebagai
        QC validated atau approved.

        Detail processing per worksheet akan
        ditambahkan pada CP4.4 melalui Processing_Log.
        """

        worksheet = workbook.create_sheet(
            self.config.control_sheet_name
        )

        controls = [
            (
                "Item",
                "Keterangan",
            ),
            (
                "Source file",
                source_file.name,
            ),
            (
                "Processing mode",
                "Multi-worksheet",
            ),
            (
                "Candidate records",
                len(records),
            ),
            (
                "Record status",
                (
                    "Candidate parser output; "
                    "not QC validation/approval"
                ),
            ),
            (
                "Worksheet traceability",
                (
                    "Source Worksheet retained "
                    "per record"
                ),
            ),
            (
                "OFF handling",
                (
                    "Tidak dibuat sebagai "
                    "record pengukuran"
                ),
            ),
            (
                "Total handling",
                (
                    "Bukan parameter; menjadi "
                    "batas pembacaan SP"
                ),
            ),
            (
                "Blank value",
                (
                    "Tidak dibuat sebagai "
                    "record pengukuran"
                ),
            ),
            (
                "STD '-'",
                (
                    "Dipertahankan sesuai source; "
                    "tidak dianggap 0"
                ),
            ),
        ]

        for row in controls:
            worksheet.append(
                row
            )

        self._format_header(
            worksheet
        )

        worksheet.column_dimensions[
            "A"
        ].width = 24

        worksheet.column_dimensions[
            "B"
        ].width = 65

    # =====================================================
    # PROCESSING LOG — CP4.4
    # =====================================================

    @staticmethod
    def _split_processing_message(
        message: str,
    ) -> tuple[str, list[str]]:
        """Pisahkan message summary dari diagnostic event yang terstruktur.

        Service mempertahankan message sebagai string untuk kompatibilitas.
        Exporter memecah diagnostic yang dikenali menjadi row EVENT agar dapat
        dibaca/filter langsung di Excel. Message non-diagnostic tetap berada
        pada row SUMMARY.
        """

        parts = [
            part.strip()
            for part in str(message or "").split(" | ")
            if part.strip()
        ]

        summary_parts: list[str] = []
        diagnostic_parts: list[str] = []

        for part in parts:
            if re.match(
                r"^(TRACE|REVIEW|WARNING)\s+[A-Z_]+\b",
                part,
            ):
                diagnostic_parts.append(part)
            else:
                summary_parts.append(part)

        return " | ".join(summary_parts), diagnostic_parts

    @staticmethod
    def _literal_text(
        value: str | None,
    ) -> str:
        """Kembalikan repr-style diagnostic value sebagai text biasa."""

        if value is None:
            return ""

        value = value.strip()
        if not value:
            return ""

        try:
            parsed = ast.literal_eval(value)
        except (ValueError, SyntaxError):
            return value.strip("'\"")

        return "" if parsed is None else str(parsed)

    @classmethod
    def _parse_processing_event(
        cls,
        event_message: str,
    ) -> dict[str, str]:
        """Parse diagnostic event yang diterbitkan parser untuk presentation."""

        parts = event_message.split(" ", 2)
        if len(parts) < 2:
            return {
                "level": "",
                "event": "",
                "location": "",
                "source_text": "",
                "canonical_identity": "",
            }

        level = parts[0].strip()
        event = parts[1].strip()
        remainder = parts[2] if len(parts) >= 3 else ""

        location_match = re.search(
            r"\bR\d+C\d+\b",
            remainder,
        )

        source_match = re.search(
            r"\bsource=(.+?)(?=\s+->\s+|$)",
            remainder,
        )

        canonical_match = re.search(
            r"\s+->\s+(.+)$",
            remainder,
        )

        return {
            "level": level,
            "event": event,
            "location": (
                location_match.group(0)
                if location_match
                else ""
            ),
            "source_text": cls._literal_text(
                source_match.group(1)
                if source_match
                else None
            ),
            "canonical_identity": cls._literal_text(
                canonical_match.group(1)
                if canonical_match
                else None
            ),
        }

    def _write_processing_log(
        self,
        workbook,
        worksheet_results,
    ) -> None:
        """Menulis summary worksheet dan diagnostic event sebagai row terpisah.

        Row SUMMARY mempertahankan status/record count existing sehingga
        monthly counting tetap menghitung satu status per worksheet. Row EVENT
        sengaja mengosongkan Status dan Record Count agar diagnostic tidak
        menggandakan worksheet count pada konsolidasi bulanan.

        Processing_Log adalah execution/observability trace, bukan QC
        validation/approval.
        """

        worksheet = workbook.create_sheet(
            "Processing_Log"
        )

        worksheet.append(
            [
                "Worksheet",
                "Visibility",
                "Status",
                "Record Count",
                "Message",
                "Entry Type",
                "Level",
                "Event",
                "Location",
                "Source Text",
                "Canonical Identity",
            ]
        )

        for result in worksheet_results:
            summary_message, diagnostic_events = (
                self._split_processing_message(
                    result.message
                )
            )

            worksheet.append(
                [
                    result.worksheet_name,
                    result.visibility,
                    result.status,
                    result.record_count,
                    summary_message,
                    "SUMMARY",
                    "",
                    "",
                    "",
                    "",
                    "",
                ]
            )

            for event_message in diagnostic_events:
                event = self._parse_processing_event(
                    event_message
                )

                worksheet.append(
                    [
                        result.worksheet_name,
                        result.visibility,
                        "",
                        "",
                        event_message,
                        "EVENT",
                        event["level"],
                        event["event"],
                        event["location"],
                        event["source_text"],
                        event["canonical_identity"],
                    ]
                )

        self._format_header(
            worksheet
        )

        worksheet.freeze_panes = "A2"

        worksheet.auto_filter.ref = (
            worksheet.dimensions
        )

        widths = {
            "A": 28,
            "B": 14,
            "C": 14,
            "D": 16,
            "E": 70,
            "F": 14,
            "G": 10,
            "H": 30,
            "I": 14,
            "J": 45,
            "K": 35,
        }

        for column, width in (
            widths.items()
        ):
            worksheet.column_dimensions[
                column
            ].width = width

        for row in worksheet.iter_rows(
            min_row=2,
            min_col=5,
            max_col=11,
        ):
            for cell in row:
                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True,
                )
