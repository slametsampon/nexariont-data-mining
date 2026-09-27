from __future__ import annotations

import argparse
import re
import subprocess
import sys
import tempfile
from collections import Counter
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


DOMAIN_ROUTING = (
    ("NPG", "NPG"),
    ("Octanol", "Octanol"),
    ("Syn Gas", "Syngas"),
    ("Utility", "Utility"),
    ("WWT", "wwt"),
)

MASTER_BASENAME = "Master-Data.xlsx"
MASTER_SOURCE_SHEET = "Sampling-Point"
CORE_MASTER_SHEET = "QA Review"

PROVENANCE_HEADERS = [
    "Source Month",
    "Source Week",
    "Domain",
    "Source Workbook",
]

ROMAN_WEEK = {
    "I": 1,
    "II": 2,
    "III": 3,
    "IV": 4,
    "V": 5,
}


def norm_name(value: str) -> str:
    return " ".join(str(value).strip().casefold().split())


def safe_name(value: str) -> str:
    cleaned = re.sub(r'[<>:"/\\|?*]+', "_", value)
    cleaned = re.sub(r"\s+", "_", cleaned.strip())
    return cleaned or "MONTH"


def week_sort_key(path: Path):
    name = path.name.strip()
    m = re.search(r"(?i)\bminggu\s+([ivx]+|\d+)\b", name)
    if m:
        token = m.group(1).upper()
        if token.isdigit():
            return (0, int(token), norm_name(name))
        if token in ROMAN_WEEK:
            return (0, ROMAN_WEEK[token], norm_name(name))
    return (1, 999, norm_name(name))


def find_week_dirs(root: Path) -> list[Path]:
    weeks = []
    for p in root.iterdir():
        if p.is_dir() and norm_name(p.name).startswith("minggu "):
            weeks.append(p)
    return sorted(weeks, key=week_sort_key)


def find_child_dir(parent: Path, expected_name: str) -> Path | None:
    expected = norm_name(expected_name)
    for p in parent.iterdir():
        if p.is_dir() and norm_name(p.name) == expected:
            return p
    return None


def find_source_workbooks(domain_dir: Path) -> list[Path]:
    files = []
    for p in domain_dir.rglob("*"):
        if not p.is_file():
            continue
        if p.name.startswith("~$"):
            continue
        if p.suffix.casefold() not in {".xls", ".xlsx"}:
            continue
        files.append(p)
    return sorted(files, key=lambda p: norm_name(str(p.relative_to(domain_dir))))


def read_sheet_rows(ws):
    return [list(row) for row in ws.iter_rows(values_only=True)]


def normalize_header(row) -> list[str]:
    return ["" if v is None else str(v).strip() for v in row]


def append_with_provenance(ws_out, provenance, rows, expected_header_state, sheet_name):
    if not rows:
        return 0

    header = normalize_header(rows[0])

    if expected_header_state["header"] is None:
        expected_header_state["header"] = header
        ws_out.append(PROVENANCE_HEADERS + header)
    elif header != expected_header_state["header"]:
        raise ValueError(
            f"{sheet_name} header mismatch. "
            f"Expected={expected_header_state['header']!r}; Actual={header!r}"
        )

    count = 0
    for row in rows[1:]:
        if not any(v is not None and str(v).strip() != "" for v in row):
            continue
        ws_out.append(list(provenance) + list(row))
        count += 1
    return count


def processing_counts(rows):
    result = Counter()
    if not rows:
        return result

    header = normalize_header(rows[0])
    index = {name.casefold(): i for i, name in enumerate(header)}
    status_idx = index.get("status")
    if status_idx is None:
        return result

    for row in rows[1:]:
        if status_idx >= len(row):
            continue
        value = row[status_idx]
        if value is None:
            continue
        result[str(value).strip().upper()] += 1
    return result


def add_control_section(ws_control, provenance, control_rows):
    ws_control.append([])
    ws_control.append([
        "Source Month", provenance[0],
        "Source Week", provenance[1],
        "Domain", provenance[2],
        "Source Workbook", provenance[3],
    ])

    if not control_rows:
        ws_control.append(["(Control sheet empty)"])
        return

    for row in control_rows:
        ws_control.append(list(row))


def set_simple_formatting(ws, freeze="A2", autofilter=True, max_width=45):
    if ws.max_row >= 1:
        for cell in ws[1]:
            cell.font = Font(bold=True)

    if freeze:
        ws.freeze_panes = freeze

    if autofilter and ws.max_row >= 1 and ws.max_column >= 1:
        ws.auto_filter.ref = ws.dimensions

    sample_rows = min(ws.max_row, 250)
    for col_idx in range(1, ws.max_column + 1):
        width = 10
        for row_idx in range(1, sample_rows + 1):
            value = ws.cell(row=row_idx, column=col_idx).value
            if value is None:
                continue
            width = max(width, min(len(str(value)) + 2, max_width))
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def tail_message(stdout: str, stderr: str, limit: int = 2000) -> str:
    text = "\n".join(
        x.strip()
        for x in (stdout or "", stderr or "")
        if x and x.strip()
    )
    return text if len(text) <= limit else text[-limit:]


def build_core_master_adapter(master_path: Path, adapter_path: Path) -> Path:
    """
    Read canonical Master-Data.xlsx / Sampling-Point and create a temporary
    adapter workbook for the released v0.2 core, which still expects the
    legacy worksheet name 'QA Review'. The canonical master is not modified.
    """
    wb_source = load_workbook(
        master_path,
        data_only=True,
        read_only=True,
    )
    try:
        if MASTER_SOURCE_SHEET not in wb_source.sheetnames:
            raise ValueError(
                f"Worksheet master '{MASTER_SOURCE_SHEET}' tidak ditemukan "
                f"dalam {master_path}"
            )

        ws_source = wb_source[MASTER_SOURCE_SHEET]

        wb_adapter = Workbook()
        ws_adapter = wb_adapter.active
        ws_adapter.title = CORE_MASTER_SHEET

        for row in ws_source.iter_rows(values_only=True):
            ws_adapter.append(list(row))

        wb_adapter.save(adapter_path)
        return adapter_path
    finally:
        wb_source.close()


def main() -> int:
    repo = Path(__file__).resolve().parent

    ap = argparse.ArgumentParser(
        description=(
            "NEXARIONT QC monthly runner. "
            "Input monthly source folder and output folder can be different locations "
            "(e.g. input on server, output on local PC)."
        )
    )
    ap.add_argument(
        "--input-root",
        required=True,
        type=Path,
        help=(
            "Monthly source folder. May be local path or UNC/server path. "
            "Example: \\\\SERVER\\QC\\8-AGUSTUS 2026"
        ),
    )
    ap.add_argument(
        "--output-dir",
        required=True,
        type=Path,
        help=(
            "Destination folder for the single consolidated monthly XLSX. "
            "May be different from input location."
        ),
    )
    ap.add_argument(
        "--master",
        type=Path,
        default=None,
        help=(
            "Optional canonical QC master path. "
            "Default: <repository>\\doc\\Master-Data.xlsx. "
            "Sampling Point data is read from worksheet 'Sampling-Point'."
        ),
    )
    ap.add_argument(
        "--output-name",
        default=None,
        help=(
            "Optional output filename. "
            "Default: QC_Data_Mining_<monthly-folder-name>.xlsx"
        ),
    )
    ap.add_argument(
        "--overwrite",
        action="store_true",
        help="Allow replacing an existing final monthly output file.",
    )
    args = ap.parse_args()

    input_root = args.input_root.expanduser()
    output_dir = args.output_dir.expanduser()

    main_py = repo / "main.py"
    if not main_py.is_file():
        raise SystemExit(f"MAIN_PY_NOT_FOUND: {main_py}")

    master = (
        args.master.expanduser()
        if args.master is not None
        else repo / "doc" / MASTER_BASENAME
    )

    if not input_root.is_dir():
        raise SystemExit(f"INPUT_ROOT_NOT_FOUND: {input_root}")

    if not master.is_file():
        raise SystemExit(
            "MASTER_NOT_FOUND:\n"
            f"  {master}\n\n"
            "Expected default location:\n"
            f"  {repo / 'doc' / MASTER_BASENAME}\n"
            "Or specify another file with --master."
        )

    weeks = find_week_dirs(input_root)
    if not weeks:
        raise SystemExit(f"NO_WEEK_FOLDERS_FOUND under: {input_root}")

    output_dir.mkdir(parents=True, exist_ok=True)

    output_name = (
        args.output_name
        or f"QC_Data_Mining_{safe_name(input_root.name)}.xlsx"
    )
    if not output_name.casefold().endswith(".xlsx"):
        output_name += ".xlsx"

    final_output = output_dir / output_name

    if final_output.exists() and not args.overwrite:
        raise SystemExit(
            f"OUTPUT_EXISTS: {final_output}\n"
            "Use --overwrite only if replacement is intentional."
        )

    wb_month = Workbook()
    ws_data = wb_month.active
    ws_data.title = "Data_Mining"
    ws_control = wb_month.create_sheet("Control")
    ws_log = wb_month.create_sheet("Processing_Log")
    ws_summary = wb_month.create_sheet("Run_Summary")

    data_header_state = {"header": None}
    log_header_state = {"header": None}

    ws_summary.append([
        "Source Month",
        "Source Week",
        "Domain",
        "Source Workbook",
        "Source Path",
        "Execution",
        "Candidate Records",
        "Worksheets",
        "Success",
        "Excluded",
        "Error",
        "Message",
    ])

    totals = Counter()
    source_counter = 0

    with tempfile.TemporaryDirectory(prefix="nexariont_qc_monthly_") as tmp:
        tmp_dir = Path(tmp)

        adapted_master = build_core_master_adapter(
            master,
            tmp_dir / "_Master-Data_core_adapter.xlsx",
        )

        for week in weeks:
            print(f"\n=== {week.name} ===")

            for folder_name, parser_domain in DOMAIN_ROUTING:
                domain_dir = find_child_dir(week, folder_name)

                if domain_dir is None:
                    print(f"[MISSING] {folder_name}")
                    ws_summary.append([
                        input_root.name,
                        week.name,
                        parser_domain,
                        "",
                        "",
                        "MISSING_DOMAIN_FOLDER",
                        0, 0, 0, 0, 0,
                        f"Folder '{folder_name}' tidak ditemukan.",
                    ])
                    totals["missing_domain_folder"] += 1
                    continue

                source_files = find_source_workbooks(domain_dir)

                if not source_files:
                    print(f"[EMPTY]   {folder_name}")
                    ws_summary.append([
                        input_root.name,
                        week.name,
                        parser_domain,
                        "",
                        str(domain_dir),
                        "NO_WORKBOOKS",
                        0, 0, 0, 0, 0,
                        "Tidak ada source .xls/.xlsx.",
                    ])
                    totals["no_workbooks"] += 1
                    continue

                print(f"[{folder_name}] {len(source_files)} workbook(s)")

                for source in source_files:
                    source_counter += 1
                    temp_output = tmp_dir / f"{source_counter:04d}.xlsx"

                    cmd = [
                        sys.executable,
                        str(main_py),
                        "--input", str(source),
                        "--output", str(temp_output),
                        "--master", str(adapted_master),
                        "--domain", parser_domain,
                    ]

                    print(f"  {source.name}")

                    proc = subprocess.run(
                        cmd,
                        cwd=repo,
                        text=True,
                        capture_output=True,
                    )

                    totals["workbooks_attempted"] += 1

                    provenance = [
                        input_root.name,
                        week.name,
                        parser_domain,
                        source.name,
                    ]

                    execution = "OK"
                    message = ""
                    candidate_records = 0
                    counts = Counter()

                    if proc.returncode != 0 or not temp_output.exists():
                        execution = "EXEC_FAILED"
                        message = tail_message(proc.stdout, proc.stderr)
                        totals["exec_failed"] += 1

                        ws_summary.append([
                            *provenance,
                            str(source),
                            execution,
                            0, 0, 0, 0, 0,
                            message,
                        ])
                        continue

                    try:
                        wb_src = load_workbook(
                            temp_output,
                            data_only=True,
                            read_only=True,
                        )
                        try:
                            required = {"Data_Mining", "Control", "Processing_Log"}
                            missing_sheets = required - set(wb_src.sheetnames)
                            if missing_sheets:
                                raise ValueError(
                                    "Missing output sheet(s): "
                                    + ", ".join(sorted(missing_sheets))
                                )

                            data_rows = read_sheet_rows(wb_src["Data_Mining"])
                            control_rows = read_sheet_rows(wb_src["Control"])
                            log_rows = read_sheet_rows(wb_src["Processing_Log"])

                            candidate_records = append_with_provenance(
                                ws_data,
                                provenance,
                                data_rows,
                                data_header_state,
                                "Data_Mining",
                            )

                            append_with_provenance(
                                ws_log,
                                provenance,
                                log_rows,
                                log_header_state,
                                "Processing_Log",
                            )

                            add_control_section(
                                ws_control,
                                provenance,
                                control_rows,
                            )

                            counts = processing_counts(log_rows)

                        finally:
                            wb_src.close()

                    except Exception as exc:
                        execution = "OUTPUT_READ_ERROR"
                        message = f"{type(exc).__name__}: {exc}"
                        totals["output_read_error"] += 1

                    worksheets = (
                        counts.get("SUCCESS", 0)
                        + counts.get("EXCLUDED", 0)
                        + counts.get("ERROR", 0)
                    )

                    totals["candidate_records"] += candidate_records
                    totals["worksheets"] += worksheets
                    totals["success"] += counts.get("SUCCESS", 0)
                    totals["excluded"] += counts.get("EXCLUDED", 0)
                    totals["error"] += counts.get("ERROR", 0)

                    ws_summary.append([
                        *provenance,
                        str(source),
                        execution,
                        candidate_records,
                        worksheets,
                        counts.get("SUCCESS", 0),
                        counts.get("EXCLUDED", 0),
                        counts.get("ERROR", 0),
                        message,
                    ])

    if data_header_state["header"] is None:
        ws_data.append(PROVENANCE_HEADERS + [
            "Source Worksheet",
            "Tanggal",
            "Jam",
            "Sampling Point (SP)",
            "Parameter",
            "Unit",
            "Min",
            "Max",
            "Nilai",
        ])

    if log_header_state["header"] is None:
        ws_log.append(PROVENANCE_HEADERS + [
            "Worksheet",
            "Visibility",
            "Status",
            "Record Count",
            "Message",
        ])

    set_simple_formatting(ws_data, freeze="A2", autofilter=True)
    set_simple_formatting(ws_log, freeze="A2", autofilter=True)
    set_simple_formatting(ws_summary, freeze="A2", autofilter=True)
    set_simple_formatting(ws_control, freeze=None, autofilter=False)

    if final_output.exists():
        final_output.unlink()

    wb_month.save(final_output)

    print("\n" + "=" * 72)
    print("NEXARIONT QC MONTHLY MINING COMPLETED")
    print("=" * 72)
    print(f"Input root            : {input_root}")
    print(f"Output directory      : {output_dir}")
    print(f"Master                : {master}")
    print(f"Master worksheet      : {MASTER_SOURCE_SHEET}")
    print(f"Month folder          : {input_root.name}")
    print(f"Weeks detected        : {len(weeks)}")
    print(f"Workbooks attempted   : {totals['workbooks_attempted']}")
    print(f"Candidate records     : {totals['candidate_records']}")
    print(f"Worksheets            : {totals['worksheets']}")
    print(f"Success               : {totals['success']}")
    print(f"Excluded              : {totals['excluded']}")
    print(f"Error                 : {totals['error']}")
    print(f"Execution failures    : {totals['exec_failed']}")
    print(f"Output read errors    : {totals['output_read_error']}")
    print(f"Missing domain folder : {totals['missing_domain_folder']}")
    print(f"Empty domain folder   : {totals['no_workbooks']}")
    print(f"Output file           : {final_output}")

    if (
        totals["error"]
        or totals["exec_failed"]
        or totals["output_read_error"]
        or totals["missing_domain_folder"]
        or totals["no_workbooks"]
    ):
        print("\nATTENTION: lihat sheet Run_Summary dan Processing_Log.")
    else:
        print("\nNo processing exception recorded.")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
